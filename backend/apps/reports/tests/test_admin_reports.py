"""Step 4.5: choir / member / participation reports and audited exports (REP-*, ADM-07)."""

import csv
import io
from datetime import timedelta

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.evaluations import services
from apps.evaluations.constants import ModerationStatus
from apps.evaluations.models import Evaluation
from apps.evaluations.tests.factories import make_criterion, make_cycle, make_form
from apps.members.tests.factories import MemberProfileFactory, voice_section
from apps.reports import selectors
from apps.reports.exports import Document, Table, safe_text, to_csv

pytestmark = pytest.mark.django_db

FORMULA = '=HYPERLINK("http://evil.example","x")'


def url(name, *args):
    return reverse(f"reports-admin:{name}", args=args)


@pytest.fixture
def world(super_admin):
    attendance = make_criterion(super_admin, name="Attendance")
    blend = make_criterion(super_admin, name="Blend", is_required=False)
    form = make_form(super_admin, criteria=[attendance, blend])
    cycle = make_cycle(super_admin, form=form, open=True)
    cv = {i.criterion_version.name: i.criterion_version_id for i in form.items.all()}
    s1, s2, s3, s4 = (
        MemberProfileFactory(voice_section=voice_section("Soprano")).user for _ in range(4)
    )
    b1 = MemberProfileFactory(voice_section=voice_section("Bass")).user

    def rate(evaluator, recipient, a, b, comment="Good work."):
        return services.submit_evaluation(
            evaluator=evaluator,
            recipient_public_id=recipient.public_id,
            ratings=[
                {"criterion_version": cv["Attendance"], "value": a},
                {"criterion_version": cv["Blend"], "value": b},
            ],
            comment=comment,
        )

    rate(s1, s2, 5, None)
    rate(s2, s1, 3, None)
    rate(s3, s1, 4, 2, comment=FORMULA)
    hidden = rate(s4, s1, 1, 4)
    Evaluation.objects.filter(pk=hidden.pk).update(moderation_status=ModerationStatus.HIDDEN)
    rate(s1, b1, 2, 5)
    return {"cycle": cycle, "s1": s1, "s2": s2, "b1": b1, "blend": blend}


@pytest.mark.parametrize("name", ["choir", "participation", "export"])
def test_only_super_admins(api_client, member_client, name):
    assert api_client.get(url(name)).status_code == 401
    assert member_client.get(url(name)).status_code == 403


def test_member_report_is_admin_only(world, member_client):
    assert member_client.get(url("member", world["s1"].public_id)).status_code == 403


def test_choir_summary(world, admin_client_api):
    data = admin_client_api.get(url("choir")).json()
    assert data["totals"] == {"evaluations": 4, "evaluators": 3, "recipients": 3}  # hidden left out
    attendance, blend = data["criteria"]
    assert attendance["criterion"] == "Attendance"
    assert (attendance["rated"], attendance["not_rated"]) == (4, 0)
    assert attendance["average"] == 3.5
    assert attendance["average_percent"] == 62.5
    # 2 ratings < threshold 3: hidden; "not rated" counted, never as zero.
    assert (blend["rated"], blend["not_rated"], blend["suppressed"]) == (2, 2, True)
    assert blend["average"] is None
    sections = {s["name"]: s for s in data["voice_sections"]}
    assert sections["Soprano"]["evaluators"] == 3
    assert sections["Soprano"]["suppressed"] is False
    assert sections["Bass"]["suppressed"] is True  # one evaluator: could be identified
    assert sections["Bass"]["average_percent"] is None
    assert data["trend"][0]["evaluations"] == 4


def test_choir_filters(world, admin_client_api):
    bass = voice_section("Bass").pk
    data = admin_client_api.get(url("choir"), {"voice_section": bass}).json()
    assert data["totals"]["evaluations"] == 1
    data = admin_client_api.get(url("choir"), {"criterion": world["blend"].public_id}).json()
    assert [c["criterion"] for c in data["criteria"]] == ["Blend"]
    tomorrow = (timezone.now() + timedelta(days=1)).date()
    data = admin_client_api.get(url("choir"), {"submitted_from": tomorrow}).json()
    assert data["totals"]["evaluations"] == 0
    bad = admin_client_api.get(
        url("choir"), {"submitted_from": tomorrow, "submitted_to": timezone.now().date()}
    )
    assert bad.status_code == 400
    assert "submitted_to" in bad.json()["error"]["fields"]


def test_member_summary(world, admin_client_api):
    data = admin_client_api.get(url("member", world["s1"].public_id)).json()
    assert (data["evaluations_received"], data["evaluations_given"]) == (2, 2)
    attendance, blend = data["criteria"]
    assert (attendance["average"], attendance["rated"], attendance["suppressed"]) == (3.5, 2, False)
    assert (blend["average"], blend["rated"], blend["not_rated"]) == (2.0, 1, 1)


def test_participation(world, admin_client_api):
    assert admin_client_api.get(url("participation")).status_code == 400
    data = admin_client_api.get(url("participation"), {"cycle": world["cycle"].public_id}).json()
    by_id = {m["public_id"]: m for m in data["members"]}
    s1 = by_id[str(world["s1"].public_id)]
    assert (s1["given"], s1["possible"], s1["received"], s1["percent"]) == (2, 4, 3, 50.0)
    assert by_id[str(world["b1"].public_id)]["given"] == 0
    assert data["totals"]["possible_pairs"] == 20
    assert data["totals"]["submitted_pairs"] == 5
    # Ordered by voice section (seeded sort order: Soprano … Bass), then last name.
    assert [m["voice_section"] for m in data["members"]][-1] == "Bass"


def _export(client, **params):
    return client.get(url("export"), params)


def test_csv_export_is_audited_and_has_no_names_in_filename(world, admin_client_api, super_admin):
    response = _export(admin_client_api, report="choir", file_type="csv")
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    disposition = response["Content-Disposition"]
    assert disposition.startswith('attachment; filename="choir-summary-')
    body = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(body)))
    assert rows[0] == ["Choir summary"]
    assert ["Attendance", "1", "1-5", "4", "0", "3.5", "62.5"] in rows
    assert any("hidden (too few)" in row for row in rows)
    log = AuditLog.objects.get(event_type="reports.export.generated")
    assert log.actor == super_admin
    assert log.metadata["report"] == "choir"
    assert log.metadata["file_type"] == "csv"
    assert log.metadata["identities"] is False


def test_evaluations_xlsx_has_identities_and_neutralises_formulas(world, admin_client_api):
    response = _export(
        admin_client_api, report="evaluations", file_type="xlsx", recipient=world["s1"].public_id
    )
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    header = next(r for r in values if r and r[0] == "Submitted (UTC)")
    assert "Evaluator" in header
    assert "Attendance" in header
    comments = [r[header.index("Comment")] for r in values[values.index(header) + 1 :]]
    assert f"'{FORMULA}" in comments
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type != "f"
    log = AuditLog.objects.get(event_type="reports.export.generated")
    assert log.metadata["identities"] is True
    assert log.metadata["filters"] == {"recipient": str(world["s1"].public_id)}
    assert log.metadata["rows"] == 3  # includes the hidden one, with its moderation status


def test_pdf_exports(world, admin_client_api):
    response = _export(
        admin_client_api, report="participation", file_type="pdf", cycle=world["cycle"].public_id
    )
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    member = _export(
        admin_client_api, report="member", file_type="pdf", member=world["s1"].public_id
    )
    assert member.status_code == 200
    filename = member["Content-Disposition"]
    assert f"member-summary-{str(world['s1'].public_id)[:8]}-" in filename
    assert world["s1"].last_name not in filename


@pytest.mark.parametrize(
    ("params", "field"),
    [
        ({"report": "evaluations", "file_type": "pdf"}, "file_type"),
        ({"report": "member", "file_type": "csv"}, "member"),
        ({"report": "participation", "file_type": "csv"}, "cycle"),
        ({"report": "secrets", "file_type": "csv"}, "report"),
        ({"report": "choir", "file_type": "docx"}, "file_type"),
    ],
)
def test_export_validation(world, admin_client_api, params, field):
    response = _export(admin_client_api, **params)
    assert response.status_code == 400
    assert field in response.json()["error"]["fields"]
    assert not AuditLog.objects.filter(event_type="reports.export.generated").exists()


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=1+1", "'=1+1"),
        ("+SUM(A1)", "'+SUM(A1)"),
        ("-2+3", "'-2+3"),
        ("@cmd", "'@cmd"),
        ("\tx", "'\tx"),
        ("plain", "plain"),
        (None, ""),
    ],
)
def test_safe_text(value, expected):
    assert safe_text(value) == expected


def test_csv_keeps_numbers_and_escapes_titles():
    doc = Document(
        title="=evil",
        filename_stem="x",
        tables=[Table(title="T", columns=["a", "b"], rows=[[-1, "=2"]])],
    )
    rows = list(csv.reader(io.StringIO(to_csv(doc).decode("utf-8-sig"))))
    assert rows[0] == ["'=evil"]
    assert ["-1", "'=2"] in rows


def test_selectors_handle_an_empty_choir(super_admin):
    data = selectors.choir_summary()
    assert data["totals"] == {"evaluations": 0, "evaluators": 0, "recipients": 0}
    assert data["criteria"] == []
    assert data["voice_sections"] == []
